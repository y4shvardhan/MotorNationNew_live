#!/usr/bin/env python3
"""
import_excel_news_reviews.py

Parses media_1791786863434_0f405cfc.xlsx ('Car News & 5-Line Reviews' sheet)
and generates:
1. data/news_and_reviews.json - full structured dataset keyed by vehicle_id and list
2. Updates database/motornation.db vehicles table with news and review columns
"""

import sys
import os
import json
import sqlite3
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = Path(r"C:\Users\Admin\.gemini\antigravity\brain\143c64f9-d960-431e-9f7f-9ad797ae07ee\.user_uploaded\media_1791786863434_0f405cfc.xlsx")
DB_PATH = BASE_DIR / "database" / "motornation.db"
OUTPUT_JSON = BASE_DIR / "data" / "news_and_reviews.json"
GALLERIES_PATH = BASE_DIR / "data" / "vehicle_galleries.json"


def clean_str(val):
    if val is None:
        return ""
    text = str(val).strip()
    return text


def parse_pros_cons(text):
    if not text:
        return []
    lines = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        # Remove bullet or dash prefixes
        clean = line.lstrip("•-–—* \t")
        if clean:
            lines.append(clean)
    return lines


def parse_review_lines(text):
    if not text:
        return []
    lines = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        lines.append(line)
    return lines


def main():
    if not EXCEL_PATH.exists():
        print(f"Error: Excel file not found at {EXCEL_PATH}")
        sys.exit(1)

    print(f"Loading Excel workbook from {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Car News & 5-Line Reviews"]

    # Load galleries if available
    galleries = {}
    if GALLERIES_PATH.exists():
        try:
            with open(GALLERIES_PATH, "r", encoding="utf-8") as f:
                galleries = json.load(f)
        except Exception as e:
            print(f"Warning: could not load galleries: {e}")

    # Load DB to verify IDs
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("SELECT rowid, vehicle_id, Year, Make, [Model / Model Group], [Trim / Variant], [Image URL], Image FROM vehicles ORDER BY rowid ASC")
    db_rows = cur.fetchall()

    print(f"Found {len(db_rows)} database vehicles.")

    # Headers at row 3
    # Row 4 onwards are data rows
    records = []
    records_by_id = {}
    records_by_key = {}

    row_count = 0
    for r in range(4, ws.max_row + 1):
        make = clean_str(ws.cell(r, 2).value)
        if not make:
            continue

        year = clean_str(ws.cell(r, 1).value) or "2026"
        model = clean_str(ws.cell(r, 3).value)
        trim = clean_str(ws.cell(r, 4).value)
        body_type = clean_str(ws.cell(r, 5).value)
        msrp_val = ws.cell(r, 6).value
        try:
            msrp = float(msrp_val) if msrp_val is not None else None
        except (ValueError, TypeError):
            msrp = None

        rating = clean_str(ws.cell(r, 7).value)
        headline = clean_str(ws.cell(r, 8).value)
        article = clean_str(ws.cell(r, 9).value)
        review_text = clean_str(ws.cell(r, 10).value)
        pros_raw = clean_str(ws.cell(r, 11).value)
        cons_raw = clean_str(ws.cell(r, 12).value)
        oem_url = clean_str(ws.cell(r, 13).value)
        cd_url = clean_str(ws.cell(r, 14).value)

        # Match with DB vehicle by index or by make/model
        db_car = None
        if row_count < len(db_rows):
            candidate = db_rows[row_count]
            c_make = clean_str(candidate[3]).lower()
            c_model = clean_str(candidate[4]).lower()
            if c_make == make.lower() and c_model == model.lower():
                db_car = candidate

        if not db_car:
            for c in db_rows:
                c_make = clean_str(c[3]).lower()
                c_model = clean_str(c[4]).lower()
                if c_make == make.lower() and c_model == model.lower():
                    db_car = c
                    break

        vehicle_id = db_car[0] if db_car else (row_count + 1)
        db_image = ""
        if db_car:
            db_image = clean_str(db_car[7]) or clean_str(db_car[6])

        # Gallery lookup
        vid_key = str(vehicle_id)
        name_key = f"{make}_{model}".lower()
        car_gallery = galleries.get(vid_key) or galleries.get(name_key) or []
        main_image = ""
        if car_gallery and len(car_gallery) > 0:
            main_image = car_gallery[0].get("url", "")
        if not main_image:
            main_image = db_image or f"/static/images/{make.lower()}_{model.lower().replace(' ', '_')}.jpg"

        pros_list = parse_pros_cons(pros_raw)
        cons_list = parse_pros_cons(cons_raw)
        review_lines = parse_review_lines(review_text)

        entry = {
            "id": vehicle_id,
            "vehicle_id": vehicle_id,
            "year": int(year) if year.isdigit() else 2026,
            "make": make,
            "model": model,
            "trim": trim,
            "body_type": body_type,
            "msrp": msrp,
            "rating": rating,
            "headline": headline,
            "article": article,
            "review": review_text,
            "review_lines": review_lines,
            "pros": pros_list,
            "cons": cons_list,
            "oem_url": oem_url,
            "cd_url": cd_url,
            "image": main_image,
            "gallery_count": len(car_gallery)
        }

        records.append(entry)
        records_by_id[str(vehicle_id)] = entry
        records_by_key[f"{make}_{model}".lower()] = entry

        row_count += 1

    print(f"Successfully processed {len(records)} vehicle news & review entries.")

    # Save to JSON
    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "count": len(records),
        "items": records,
        "by_id": records_by_id,
        "by_key": records_by_key
    }
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
    print(f"Saved {OUTPUT_JSON} ({os.path.getsize(OUTPUT_JSON):,} bytes)")

    # Update SQLite database
    print("Updating SQLite database columns...")
    cursor = conn.cursor()

    # Check existing columns
    cursor.execute("PRAGMA table_info(vehicles)")
    col_names = [col[1] for col in cursor.fetchall()]

    new_cols = [
        ("news_headline", "TEXT"),
        ("news_article", "TEXT"),
        ("road_test_review", "TEXT"),
        ("oem_newsroom_url", "TEXT")
    ]
    for col_name, col_type in new_cols:
        if col_name not in col_names:
            cursor.execute(f"ALTER TABLE vehicles ADD COLUMN {col_name} {col_type}")
            print(f"Added column {col_name} to vehicles table.")

    # Batch update
    update_count = 0
    for item in records:
        vid = item["vehicle_id"]
        cursor.execute("""
            UPDATE vehicles
            SET news_headline = ?,
                news_article = ?,
                road_test_review = ?,
                oem_newsroom_url = ?
            WHERE rowid = ? OR vehicle_id = ?
        """, (
            item["headline"],
            item["article"],
            item["review"],
            item["oem_url"],
            vid,
            vid
        ))
        update_count += 1

    conn.commit()
    conn.close()
    print(f"Database update complete: {update_count} rows updated.")


if __name__ == "__main__":
    main()
